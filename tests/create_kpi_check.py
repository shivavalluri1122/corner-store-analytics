"""KPI validation script (not a pytest test). Independently recomputes the
five dashboard KPIs straight from the raw synthetic CSV with plain pandas -
no pipeline or dashboard code - then compares them with the values the
application's DuckDB/analytics layer returns, and writes tests/kpi_check.xlsx.
It then prints the synthetic-data pattern sanity checks (weekend lift,
spikes, summer beverage lift, stockout gaps).

Usage:
    python tests/create_kpi_check.py [path_to_csv]
Defaults to data/synthetic/large_clean.csv, which is regenerated with the
project's generator if it is missing (it is gitignored).
Exits non-zero if any KPI does not PASS.
"""

import datetime as dt
import random
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent.parent))

from data.synthetic.generate_dataset import (  # noqa: E402
    FIRST_OF_MONTH_DAYS,
    START_DATE,
    SUMMER_MONTHS,
    build_stockout_windows,
    build_sku_weights,
)
from data.synthetic.generate_dataset import main as generate_datasets  # noqa: E402
from data.synthetic.sku_catalog import build_sku_catalog  # noqa: E402
from pipeline import analytics  # noqa: E402
from pipeline.db import build_and_load, get_connection  # noqa: E402
from pipeline.ingest import ingest_csv  # noqa: E402
from pipeline.transform import transform  # noqa: E402

DATA_DIR = Path(__file__).parent.parent / "data" / "synthetic"
OUTPUT_PATH = Path(__file__).parent / "kpi_check.xlsx"
LARGE_CSV = DATA_DIR / "large_clean.csv"
EXPECTED_LARGE_ROWS = 50_000
LARGE_ROWS_TOLERANCE = 0.10

# (display name, is_currency)
KPI_FIELDS = [
    ("Net Sales", True),
    ("Transactions", False),
    ("Average Basket", True),
    ("Units", False),
    ("Distinct Items Sold", False),
]

CENT = Decimal("0.01")


def ensure_large_dataset() -> Path:
    """Return large_clean.csv, regenerating it via the project generator if missing."""
    if not LARGE_CSV.exists():
        print(f"{LARGE_CSV} not found - regenerating with data.synthetic.generate_dataset")
        generate_datasets()
    if not LARGE_CSV.exists():
        raise SystemExit(f"Generator did not produce {LARGE_CSV}")

    rows = len(pd.read_csv(LARGE_CSV, dtype=str, usecols=["Transaction ID"]))
    low = EXPECTED_LARGE_ROWS * (1 - LARGE_ROWS_TOLERANCE)
    high = EXPECTED_LARGE_ROWS * (1 + LARGE_ROWS_TOLERANCE)
    if not low <= rows <= high:
        raise SystemExit(f"{LARGE_CSV} has {rows:,} rows; expected approximately {EXPECTED_LARGE_ROWS:,}")
    print(f"Verified {LARGE_CSV.name}: {rows:,} rows")
    return LARGE_CSV


def pick_dataset() -> Path:
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    return ensure_large_dataset()


def _to_cents(value: str) -> int:
    """'$12.98' / '-$1.50' -> integer cents, without float arithmetic."""
    text = value.strip()
    negative = text.startswith("-")
    amount = Decimal(text.replace("-", "").replace("$", "").replace(",", ""))
    cents = int(amount * 100)
    return -cents if negative else cents


def _round_cents(value: float) -> Decimal:
    return Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)


def read_raw_dates(csv_path: Path) -> tuple[dt.date, dt.date]:
    """Min/max transaction date in the raw CSV."""
    dates = pd.to_datetime(pd.read_csv(csv_path, dtype=str, usecols=["Date"])["Date"], format="%m/%d/%Y")
    return dates.min().date(), dates.max().date()


def independent_kpis(csv_path: Path, start: dt.date, end: dt.date) -> dict[str, float]:
    """Compute the five KPIs from the raw CSV using only pandas/stdlib."""
    raw = pd.read_csv(csv_path, dtype=str, usecols=["Date", "Item", "Qty", "Net Sales", "Transaction ID"])
    day = pd.to_datetime(raw["Date"], format="%m/%d/%Y").dt.date
    raw = raw[(day >= start) & (day <= end)]

    net_sales = sum(_to_cents(v) for v in raw["Net Sales"]) / 100
    transactions = int(raw["Transaction ID"].nunique())
    units = int(sum(int(q) for q in raw["Qty"]))
    distinct_items = int(raw["Item"].nunique())
    avg_basket = net_sales / transactions if transactions else 0.0
    return {
        "Net Sales": net_sales,
        "Transactions": transactions,
        "Average Basket": avg_basket,
        "Units": units,
        "Distinct Items Sold": distinct_items,
    }


def application_kpis(csv_path: Path, start: dt.date, end: dt.date) -> dict[str, float]:
    """KPIs as the dashboard computes them: ingest -> DuckDB -> analytics.compute_kpis."""
    conn = get_connection()
    build_and_load(conn, ingest_csv(csv_path))
    kpis = analytics.compute_kpis(conn, start, end)
    conn.close()
    return {
        "Net Sales": kpis.net_sales,
        "Transactions": kpis.transactions,
        "Average Basket": kpis.avg_basket,
        "Units": kpis.units,
        "Distinct Items Sold": kpis.distinct_items,
    }


def date_ranges(csv_path: Path) -> list[tuple[str, dt.date, dt.date]]:
    """Full range plus a deterministic middle-third sub-range to exercise the date filter."""
    min_date, max_date = read_raw_dates(csv_path)
    span = (max_date - min_date).days
    sub_start = min_date + dt.timedelta(days=span // 3)
    sub_end = min_date + dt.timedelta(days=2 * span // 3)
    return [("Full range", min_date, max_date), ("Middle third (date-filter check)", sub_start, sub_end)]


def compare(csv_path: Path) -> list[dict]:
    """One row per (date range, KPI). Currency compared to the cent, counts exactly."""
    rows = []
    for label, start, end in date_ranges(csv_path):
        independent = independent_kpis(csv_path, start, end)
        application = application_kpis(csv_path, start, end)
        for name, is_currency in KPI_FIELDS:
            ind, app = independent[name], application[name]
            if is_currency:
                difference = _round_cents(app) - _round_cents(ind)
                passed = difference == 0
                ind_out, app_out, diff_out = float(_round_cents(ind)), float(_round_cents(app)), float(difference)
            else:
                diff_out = app - ind
                passed = diff_out == 0
                ind_out, app_out = ind, app
            rows.append(
                {
                    "KPI": name,
                    "Independent Value": ind_out,
                    "Application/Query Value": app_out,
                    "Difference": diff_out,
                    "Status": "PASS" if passed else "FAIL",
                    "Date Range": f"{label}: {start} to {end}",
                    "is_currency": is_currency,
                }
            )
    return rows


def write_workbook(rows: list[dict], csv_path: Path, output_path: Path = OUTPUT_PATH) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI Check"
    headers = ["KPI", "Independent Value", "Application/Query Value", "Difference", "Status", "Date Range"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    for row in rows:
        ws.append([row[h] for h in headers])
        r = ws.max_row
        number_format = '"$"#,##0.00' if row["is_currency"] else "#,##0"
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).number_format = number_format
        ws.cell(row=r, column=5).fill = green if row["Status"] == "PASS" else red

    for idx, header in enumerate(headers, start=1):
        width = max(len(header), *(len(str(r[header])) for r in rows)) + 2
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"

    notes = wb.create_sheet("Methodology")
    all_pass = all(r["Status"] == "PASS" for r in rows)
    lines = [
        ("Source CSV", str(csv_path)),
        ("Validation date", dt.datetime.now().isoformat(timespec="seconds")),
        ("Overall result", "PASS" if all_pass else "FAIL"),
        ("Independent method", "Raw CSV read with pandas (all columns as text); no pipeline or dashboard code used."),
        ("Application method", "ingest_csv -> build_and_load (DuckDB) -> pipeline.analytics.compute_kpis"),
        ("Net Sales", "SUM(Net Sales) over rows in the date range (summed as integer cents)"),
        ("Transactions", "COUNT(DISTINCT Transaction ID)"),
        ("Average Basket", "Net Sales / Transactions (0 when there are no transactions)"),
        ("Units", "SUM(Qty)"),
        ("Distinct Items Sold", "COUNT(DISTINCT Item)"),
        ("Tolerance - currency", "Both values rounded to the nearest cent (half-up); PASS only if equal"),
        ("Tolerance - counts/units", "Exact match"),
        ("Date filter", "Inclusive of both start and end dates"),
        ("Reproduce", "python tests/create_kpi_check.py"),
    ]
    notes.append(["Item", "Detail"])
    for cell in notes[1]:
        cell.font = Font(bold=True)
    for line in lines:
        notes.append(list(line))
    notes.column_dimensions["A"].width = 26
    notes.column_dimensions["B"].width = 100
    for row in notes.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(output_path)


def print_kpi_report(rows: list[dict]) -> None:
    print("\n--- KPI validation (independent vs application) ---")
    for row in rows:
        fmt = "{:,.2f}" if row["is_currency"] else "{:,.0f}"
        print(
            f"[{row['Status']}] {row['KPI']:<20} independent={fmt.format(row['Independent Value'])} "
            f"application={fmt.format(row['Application/Query Value'])} "
            f"diff={fmt.format(row['Difference'])}  ({row['Date Range']})"
        )


def print_pattern_checks(csv_path: Path) -> None:
    df = ingest_csv(csv_path)
    tables = transform(df)
    line_items = tables["line_items"]
    daily = tables["daily_item_sales"]
    transactions = tables["transactions"]

    print("\n--- Dataset summary ---")
    print(f"Date range: {daily['date'].min().date()} to {daily['date'].max().date()}")
    print(f"Total transactions: {len(transactions):,}")
    print(f"Total units sold: {int(line_items['qty'].sum()):,}")
    print(f"Total net sales: ${line_items['net_sales'].sum():,.2f}")
    print(f"Avg basket size (items): {transactions['num_line_items'].mean():.2f}")
    print(f"Avg basket value: ${transactions['total_net_sales'].mean():.2f}")

    print("\n--- Category mix (% of net sales) ---")
    mix = (
        line_items.groupby("category")["net_sales"].sum().sort_values(ascending=False)
        / line_items["net_sales"].sum()
        * 100
    )
    print(mix.round(1).to_string())

    print("\n--- Top 5 SKUs by units ---")
    print(daily.groupby(["sku", "item"])["units_sold"].sum().sort_values(ascending=False).head(5))

    print("\n--- Weekend vs weekday avg daily net sales ---")
    daily_totals = line_items.groupby("date")["net_sales"].sum().reset_index()
    daily_totals["weekday"] = daily_totals["date"].dt.day_name()
    is_weekend = daily_totals["weekday"].isin(["Friday", "Saturday", "Sunday"])
    print(f"Weekend avg: ${daily_totals.loc[is_weekend, 'net_sales'].mean():.2f}")
    print(f"Weekday avg: ${daily_totals.loc[~is_weekend, 'net_sales'].mean():.2f}")

    print("\n--- First-of-month spike check ---")
    is_first_of_month = daily_totals["date"].dt.day.isin(FIRST_OF_MONTH_DAYS)
    print(f"Days 1-3 avg: ${daily_totals.loc[is_first_of_month, 'net_sales'].mean():.2f}")
    print(f"Other days avg: ${daily_totals.loc[~is_first_of_month, 'net_sales'].mean():.2f}")

    print("\n--- Summer beverage lift check ---")
    bev = line_items[line_items["category"] == "Beverages"].copy()
    bev["month"] = bev["date"].dt.month
    is_summer = bev["month"].isin(SUMMER_MONTHS)
    summer_daily_avg = bev.loc[is_summer].groupby("date")["net_sales"].sum().mean()
    other_daily_avg = bev.loc[~is_summer].groupby("date")["net_sales"].sum().mean()
    print(f"Summer months avg daily beverage sales: ${summer_daily_avg:.2f}")
    print(f"Other months avg daily beverage sales: ${other_daily_avg:.2f}")

    print("\n--- Stockout gap check (5 SKUs) ---")
    catalog = build_sku_catalog()
    rng = random.Random(42)
    build_sku_weights(catalog, rng)  # must match generate_line_items' RNG consumption order
    windows = build_stockout_windows(catalog, rng)
    for sku, (start_idx, end_idx) in windows.items():
        window_start = pd.Timestamp(START_DATE) + pd.Timedelta(days=start_idx)
        window_end = pd.Timestamp(START_DATE) + pd.Timedelta(days=end_idx)
        sales_in_window = daily[
            (daily["sku"] == sku) & (daily["date"] >= window_start) & (daily["date"] <= window_end)
        ]
        print(
            f"{sku}: window {window_start.date()} to {window_end.date()} "
            f"({(end_idx - start_idx) + 1} days) -> {len(sales_in_window)} sale-days recorded (expect 0)"
        )


def main() -> int:
    csv_path = pick_dataset()
    print(f"Checking: {csv_path}")

    rows = compare(csv_path)
    write_workbook(rows, csv_path)
    print_kpi_report(rows)
    print(f"\nWrote {OUTPUT_PATH}")

    print_pattern_checks(csv_path)

    failed = [r for r in rows if r["Status"] != "PASS"]
    if failed:
        print(f"\nKPI CHECK FAILED: {len(failed)} mismatch(es)")
        return 1
    print("\nKPI CHECK PASSED: all KPIs match")
    return 0


if __name__ == "__main__":
    sys.exit(main())
